"""All Excel (OpenPyXL) reading, writing, formatting and safe-save logic."""

from __future__ import annotations

import logging
import os
import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

import config

logger = logging.getLogger(__name__)


class ExcelHandlerError(Exception):
    """Base class for all workbook-related errors."""


class FolderAccessError(ExcelHandlerError):
    """The selected folder does not exist or cannot be used."""


class WorkbookLockedError(ExcelHandlerError):
    """The workbook is open in another program or access was denied."""


class WorkbookCorruptError(ExcelHandlerError):
    """The workbook exists but cannot be read as a valid .xlsx file."""


class InvalidWorkbookError(ExcelHandlerError):
    """The workbook is readable but lacks the expected sheet or headers."""


class WorkbookSaveError(ExcelHandlerError):
    """Saving the workbook failed."""


@dataclass(frozen=True)
class RawRow:
    """One data row exactly as found in the worksheet (values as text)."""

    row_number: int
    case_id: str
    explanation: str
    created_at: str


def _cell_text(value: Any) -> str:
    """Convert a raw cell value to text."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime(config.TIMESTAMP_FORMAT)
    return str(value)


class ExcelHandler:
    """Creates, validates, reads and appends to ``cases.xlsx`` in a folder."""

    def __init__(self, folder: str | Path) -> None:
        self.folder = Path(folder)
        self.path = self.folder / config.WORKBOOK_NAME

    # ------------------------------------------------------------------ setup
    def initialize(self) -> bool:
        """Make sure a valid workbook exists.

        Returns True if a new workbook was created, False if an existing,
        valid workbook was found. An existing workbook is never overwritten;
        if it is invalid an exception is raised and the file is left alone.
        """
        self._check_folder()
        if self.path.exists():
            self._open_workbook().close()
            logger.info("Using existing workbook %s", self.path)
            return False

        wb = Workbook()
        ws = wb.active
        ws.title = config.WORKSHEET_NAME
        ws.append(list(config.HEADERS))
        self._apply_formatting(ws, last_row=1)
        try:
            self._safe_save(wb)
        finally:
            wb.close()
        logger.info("Created new workbook %s", self.path)
        return True

    # ---------------------------------------------------------------- reading
    def read_rows(self) -> list[RawRow]:
        """Return all non-empty data rows from the Cases worksheet."""
        wb = self._open_workbook()
        try:
            ws = wb[config.WORKSHEET_NAME]
            rows: list[RawRow] = []
            for number, values in enumerate(
                ws.iter_rows(min_row=2, max_col=3, values_only=True), start=2
            ):
                cells = (list(values) + [None] * 3)[:3]
                if self._is_blank_row(cells):
                    continue
                rows.append(RawRow(number, *(_cell_text(v) for v in cells)))
            return rows
        finally:
            wb.close()

    # ---------------------------------------------------------------- writing
    def append_record(self, case_id: str, explanation: str, created_at: datetime) -> None:
        """Append one record and save safely. Raises on any failure."""
        wb = self._open_workbook()
        try:
            ws = wb[config.WORKSHEET_NAME]
            self._backup()
            row = self._last_data_row(ws) + 1
            ws.cell(row=row, column=1, value=case_id)
            explanation_cell = ws.cell(row=row, column=2, value=explanation)
            # Text such as "=1+1" must stay text, never become a formula.
            explanation_cell.data_type = "s"
            ws.cell(row=row, column=3, value=created_at)
            self._apply_formatting(ws, last_row=row)
            self._safe_save(wb)
        finally:
            wb.close()

        # Confirm the record really reached the file before reporting success.
        if not any(r.case_id.strip().upper() == case_id for r in self.read_rows()):
            raise WorkbookSaveError(
                f"The workbook was saved but {case_id} could not be found afterwards."
            )
        logger.info("Saved %s to %s", case_id, self.path)

    # --------------------------------------------------------------- internals
    def _check_folder(self) -> None:
        if not self.folder.exists() or not self.folder.is_dir():
            raise FolderAccessError(f"The folder '{self.folder}' does not exist.")
        if not os.access(self.folder, os.R_OK | os.W_OK | os.X_OK):
            raise FolderAccessError(
                f"The folder '{self.folder}' cannot be read or written. "
                "Choose a different folder or check its permissions."
            )

    def _open_workbook(self) -> Workbook:
        """Load and validate the workbook; never modifies the file."""
        try:
            wb = load_workbook(self.path)
        except PermissionError as exc:
            raise WorkbookLockedError(
                f"Access to '{self.path.name}' was denied. Close it in Excel "
                "(or any other program) and try again."
            ) from exc
        except FileNotFoundError as exc:
            raise InvalidWorkbookError(f"'{self.path.name}' was not found.") from exc
        except OSError as exc:
            raise ExcelHandlerError(f"Could not read '{self.path.name}': {exc}") from exc
        except (BadZipFile, InvalidFileException, KeyError, ValueError, EOFError) as exc:
            raise WorkbookCorruptError(self._corrupt_message()) from exc
        except Exception as exc:  # malformed XML etc.; the file is left untouched
            logger.exception("Unexpected error opening workbook")
            raise WorkbookCorruptError(self._corrupt_message()) from exc

        if config.WORKSHEET_NAME not in wb.sheetnames:
            wb.close()
            raise InvalidWorkbookError(
                f"'{self.path.name}' has no worksheet named '{config.WORKSHEET_NAME}'. "
                "The file was not changed."
            )
        ws = wb[config.WORKSHEET_NAME]
        found = tuple(
            _cell_text(ws.cell(row=1, column=col).value).strip()
            for col in range(1, len(config.HEADERS) + 1)
        )
        if found != config.HEADERS:
            wb.close()
            raise InvalidWorkbookError(
                f"The '{config.WORKSHEET_NAME}' sheet must start with the headers "
                f"{', '.join(config.HEADERS)}. The file was not changed."
            )
        return wb

    def _corrupt_message(self) -> str:
        return (
            f"'{self.path.name}' is corrupt or is not a valid Excel workbook. "
            "It was not changed. Repair or move it, then select the folder again."
        )

    @staticmethod
    def _is_blank_row(cells: list[Any]) -> bool:
        return all(c is None or (isinstance(c, str) and not c.strip()) for c in cells)

    def _last_data_row(self, ws: Worksheet) -> int:
        row = ws.max_row
        while row > 1 and self._is_blank_row(
            [ws.cell(row=row, column=c).value for c in (1, 2, 3)]
        ):
            row -= 1
        return row

    def _backup(self) -> None:
        """Copy the current workbook to ``cases.xlsx.bak`` (best effort)."""
        backup = self.path.with_name(self.path.name + config.BACKUP_SUFFIX)
        try:
            shutil.copy2(self.path, backup)
        except OSError as exc:
            logger.warning("Could not create backup %s: %s", backup, exc)

    def _safe_save(self, wb: Workbook) -> None:
        """Write to a temporary file, then atomically replace the real one."""
        temp = self.path.with_name(self.path.name + config.TEMP_SUFFIX)
        try:
            wb.save(temp)
            os.replace(temp, self.path)
        except PermissionError as exc:
            self._remove_temp(temp)
            raise WorkbookLockedError(
                f"Could not write to '{self.path.name}'. It may be open in Excel "
                "or the folder may be read-only. Close it and try again."
            ) from exc
        except Exception as exc:
            self._remove_temp(temp)
            logger.exception("Workbook save failed")
            raise WorkbookSaveError(f"Saving '{self.path.name}' failed: {exc}") from exc

    @staticmethod
    def _remove_temp(temp: Path) -> None:
        try:
            temp.unlink(missing_ok=True)
        except OSError:
            logger.warning("Could not remove temporary file %s", temp)

    @staticmethod
    def _apply_formatting(ws: Worksheet, last_row: int) -> None:
        """Style the header, set widths, freeze the header and enable filters."""
        header_fill = PatternFill("solid", start_color=config.HEADER_FILL_COLOR)
        header_font = Font(bold=True, color=config.HEADER_FONT_COLOR)
        for col in range(1, len(config.HEADERS) + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")
        for letter, width in config.COLUMN_WIDTHS.items():
            ws.column_dimensions[letter].width = width
        if last_row > 1:
            ws.cell(row=last_row, column=1).alignment = Alignment(vertical="top")
            ws.cell(row=last_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
            created = ws.cell(row=last_row, column=3)
            created.alignment = Alignment(vertical="top")
            created.number_format = config.EXCEL_TIMESTAMP_FORMAT
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:C{max(last_row, 1)}"
