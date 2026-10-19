from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

import config
import excel_handler
from excel_handler import (
    ExcelHandler,
    FolderAccessError,
    InvalidWorkbookError,
    WorkbookCorruptError,
    WorkbookLockedError,
    WorkbookSaveError,
)

NOW = datetime(2026, 1, 15, 9, 30, 0)


def test_initialize_creates_formatted_workbook(tmp_path):
    handler = ExcelHandler(tmp_path)
    assert handler.initialize() is True
    wb = load_workbook(handler.path)
    ws = wb[config.WORKSHEET_NAME]
    assert [c.value for c in ws[1]] == list(config.HEADERS)
    assert ws["A1"].font.bold is True
    assert ws["A1"].fill.start_color.rgb.endswith(config.HEADER_FILL_COLOR)
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:C1"
    assert ws.column_dimensions["B"].width == config.COLUMN_WIDTHS["B"]


def test_initialize_never_overwrites_existing(handler, seed_rows):
    seed_rows([("CASE-0001", "keep me", "2026-01-01 10:00:00")])
    again = ExcelHandler(handler.folder)
    assert again.initialize() is False
    assert [r.explanation for r in again.read_rows()] == ["keep me"]


def test_missing_folder_raises(tmp_path):
    with pytest.raises(FolderAccessError):
        ExcelHandler(tmp_path / "does-not-exist").initialize()


def test_corrupt_workbook_is_reported_and_untouched(tmp_path):
    path = tmp_path / config.WORKBOOK_NAME
    path.write_bytes(b"this is not a zip file")
    with pytest.raises(WorkbookCorruptError):
        ExcelHandler(tmp_path).initialize()
    assert path.read_bytes() == b"this is not a zip file"


def test_missing_worksheet_is_reported_and_untouched(tmp_path):
    wb = Workbook()
    wb.active.title = "Other"
    wb.save(tmp_path / config.WORKBOOK_NAME)
    with pytest.raises(InvalidWorkbookError):
        ExcelHandler(tmp_path).initialize()
    assert load_workbook(tmp_path / config.WORKBOOK_NAME).sheetnames == ["Other"]


def test_wrong_headers_are_reported(tmp_path):
    wb = Workbook()
    wb.active.title = config.WORKSHEET_NAME
    wb.active.append(["ID", "Text", "Date"])
    wb.save(tmp_path / config.WORKBOOK_NAME)
    with pytest.raises(InvalidWorkbookError):
        ExcelHandler(tmp_path).initialize()


def test_append_creates_backup_of_existing_workbook(handler):
    handler.append_record("CASE-0001", "first", NOW)
    backup = handler.path.with_name(handler.path.name + config.BACKUP_SUFFIX)
    assert backup.exists()


def test_unicode_and_long_text_round_trip(handler):
    text = "Ünïcödé – 顧客 ग्राहक 🙂 " * 400
    handler.append_record("CASE-0001", text, NOW)
    rows = handler.read_rows()
    assert rows[0].explanation == text
    assert rows[0].created_at == "2026-01-15 09:30:00"


def test_formula_like_text_stays_text(handler):
    handler.append_record("CASE-0001", "=1+1", NOW)
    ws = load_workbook(handler.path)[config.WORKSHEET_NAME]
    assert ws["B2"].value == "=1+1"
    assert ws["B2"].data_type == "s"


def test_filter_range_grows_with_data(handler):
    handler.append_record("CASE-0001", "a", NOW)
    handler.append_record("CASE-0002", "b", NOW)
    ws = load_workbook(handler.path)[config.WORKSHEET_NAME]
    assert ws.auto_filter.ref == "A1:C3"


def test_locked_workbook_raises_and_leaves_no_temp_file(handler, monkeypatch):
    def deny(*_args, **_kwargs):
        raise PermissionError("locked by Excel")

    monkeypatch.setattr(excel_handler.os, "replace", deny)
    with pytest.raises(WorkbookLockedError):
        handler.append_record("CASE-0001", "x", NOW)
    assert handler.read_rows() == []
    assert not list(handler.folder.glob("*" + config.TEMP_SUFFIX))


def test_generic_save_failure_raises_save_error(handler, monkeypatch):
    def fail(*_args, **_kwargs):
        raise OSError(28, "No space left on device")

    monkeypatch.setattr(excel_handler.os, "replace", fail)
    with pytest.raises(WorkbookSaveError):
        handler.append_record("CASE-0001", "x", NOW)
    assert handler.read_rows() == []
